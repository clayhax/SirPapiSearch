#!/usr/bin/python3
import argparse
import csv
import os
import re
import time
import hashlib
import threading
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections.abc import Callable
from dataclasses import dataclass, asdict
from io import BytesIO
from urllib.parse import urlparse, unquote, parse_qs, urljoin
try:
    from bs4 import BeautifulSoup
except Exception:
    BeautifulSoup = None
import requests
import zipfile
import logging
from serpapi import search
from serpapi.exceptions import HTTPError as SerpAPIHTTPError

# ---------------- Terminal Colors ----------------
CYAN = "\033[96m"
WHITE = "\033[97m"
GREEN = "\033[92m"
YELLOW = "\033[93m"
RED = "\033[91m"
DIM = "\033[90m"
RESET = "\033[0m"

def info(message):
    print(f"{CYAN}[+]{RESET} {message}")

def success(message):
    print(f"{GREEN}[✓]{RESET} {message}")

def notice(message):
    print(f"{DIM}[*]{RESET} {message}")

def warning(message):
    print(f"{YELLOW}[!]{RESET} {message}")

def error(message):
    print(f"{RED}[-]{RESET} {message}")


class SearchInterrupted(KeyboardInterrupt):
    """Ctrl-C during a search, carrying the results collected so far."""

    def __init__(self, partial_results):
        super().__init__()
        self.partial_results = partial_results

    
def serp_search_with_retry(
    params: dict,
    context: str,
    max_retries: int = 3,
):
    """
    Run a SerpAPI search with retry handling for transient HTTP errors.

    Returns:
        dict on success
        None after all retry attempts fail
    """

    for attempt in range(1, max_retries + 1):
        try:
            return search(params)

        except SerpAPIHTTPError as e:
            if attempt == max_retries:
                warning(
                    f"{context} SerpAPI request failed after "
                    f"{max_retries} attempts: {e}"
                )
                return None

            wait_s = 3 * attempt

            warning(
                f"{context} SerpAPI request failed. "
                f"Retrying in {wait_s}s "
                f"({attempt}/{max_retries})..."
            )

            time.sleep(wait_s)

    return None
    
def print_banner():
    banner = r"""
 ____  _      ____             _ ____                      _     
/ ___|(_)_ __|  _ \ __ _ _ __ (_) ___|  ___  __ _ _ __ ___| |__  
\___ \| | '__| |_) / _` | '_ \| \___ \ / _ \/ _` | '__/ __| '_ \ 
 ___) | | |  |  __/ (_| | |_) | |___) |  __/ (_| | | | (__| | | |
|____/|_|_|  |_|   \__,_| .__/|_|____/ \___|\__,_|_|  \___|_| |_|
                        |_|                                      

"""
    print(CYAN + banner + RESET)
    print(WHITE + "        SirPapiSearch v3.3 | by cl4yh4x" + RESET)
    print()

# ---------------- SerpAPI Key Configuration ----------------
# API key resolution priority:
#   1) --api-key argument
#   2) SERPAPI_KEY environment variable
#   3) HARDCODED_SERPAPI_KEY (convenient fallback; leave "" to disable)
HARDCODED_SERPAPI_KEY = ""  # e.g. "your_serpapi_key_here"

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None
logging.getLogger("pypdf").setLevel(logging.ERROR)

try:
    from docx import Document
except Exception:
    Document = None

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

try:
    import olefile
except Exception:
    olefile = None


# ---- Heuristics / regex ----
INTERNAL_PATH_PATTERNS = [
    r"[A-Za-z]:\\",            # C:\...
    r"\\\\[A-Za-z0-9_.-]+\\",  # \\server\share...
    r"/Users/",                # macOS
    r"/home/",                 # Linux
]
_internal_path_re = re.compile("|".join(INTERNAL_PATH_PATTERNS))

_email_re = re.compile(r"\b[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[A-Za-z]{2,}\b")
# DOMAIN\username tokens only. The lookbehind skips path segments such as
# C:\Users\bob or \\server\share, which are reported as internal paths.
_user_re = re.compile(r"(?<![\w.\\:/-])[A-Za-z0-9_.-]{2,}\\[A-Za-z0-9_.$-]{2,}\b")

_linkedin_profile_url_re = re.compile(
    r"https?://(?:[a-z0-9-]+\.)*linkedin\.com/in/[^&?#\s]+",
    flags=re.IGNORECASE,
)

# ZoomInfo person pages: https://www.zoominfo.com/p/John-Smith/1234567890
_zoominfo_profile_url_re = re.compile(
    r"https?://(?:[a-z0-9-]+\.)*zoominfo\.com/p/[^&?#\s]+",
    flags=re.IGNORECASE,
)

KEYWORDS = [
    "password", "passwd", "pwd",
    "token", "apikey", "api_key", "secret", "client_secret",
    "authorization", "bearer",
    "private key", "ssh-rsa", "BEGIN PRIVATE KEY", "BEGIN RSA PRIVATE KEY",
    "connectionstring", "jdbc:", "odbc", "ldap", "saml",
]

DOCUMENT_PATHS = [
    "/documents",
    "/docs",
    "/forms",
    "/resources",
    "/downloads",
    "/policies",
    "/procedures",
    "/board",
    "/boards",
    "/departments",
    "/staff",
    "/media",
    "/files",
    "/uploads",
    "/public",
    "/publications",
    "/public-notices",
    "/notices",
    "/agendas",
    "/minutes",
    "/meetings",
    "/records",
    "/reports",
    "/archive",
    "/archives",
    "/backup",
    "/data",
    "/assets"
]


# ---------------- Employee Sources (SerpAPI Google results only) ----------------
HONORIFICS = {
    "mr", "mrs", "ms", "miss", "mx", "dr", "prof", "sir", "madam", "dame",
}
SUFFIXES = {
    "jr", "sr", "ii", "iii", "iv", "v",
    "md", "phd", "dds", "dvm", "esq", "mba", "pe", "cissp",
}

PROFESSIONAL_CREDENTIALS = {
    "cic", "cpia", "cpcu", "crm", "arm", "ains",
    "cpa", "pmp", "rn", "msn", "np", "jd",
    "clu", "chfc", "cfp", "cfa", "caia", "aif",
    "sphr", "phr", "ccws", "cisr", 
     "phd", "ches", "ms", "rd", "ldn", "cdces"     
}

TRAILING_NAME_TOKENS = SUFFIXES | PROFESSIONAL_CREDENTIALS

LASTNAME_PARTICLES = {
    "da", "de", "del", "della", "der", "di", "du", "la", "le", "los", "las",
    "van", "von", "st", "st.", "san", "santa",
}


def strip_accents(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", s)
        if not unicodedata.combining(c)
    )


def normalize_name_token(s: str) -> str:
    s = strip_accents(s)
    s = s.replace("’", "'").replace("`", "'")
    s = re.sub(r"[^\w\-\'.]", "", s, flags=re.UNICODE)
    return s


def clean_linkedin_title_to_name(title: str) -> str:
    """
    Handles common Google result title shapes, e.g.:
      'John Doe - CompanyName'
      'John Doe - Board Chair, CompanyName'
      'LinkedIn - John Doe'
      'LinkedIn · John Doe'
      'John Doe - CompanyName | LinkedIn'
    """
    if not title:
        return ""

    t = title.strip()

    # Strip trailing branding
    t = re.sub(r"\s*\|\s*LinkedIn\s*$", "", t, flags=re.IGNORECASE)

    # Strip leading "LinkedIn - " or "LinkedIn · " or "LinkedIn:"
    t = re.sub(r"^\s*LinkedIn\s*[-·:]\s*", "", t, flags=re.IGNORECASE)

    # Split on dash variants and take first segment as candidate name
    t = re.split(r"\s+[-–—]\s+", t, maxsplit=1)[0].strip()

    # Collapse whitespace
    t = re.sub(r"\s{2,}", " ", t).strip()

    return t

def clean_zoominfo_title_to_name(title: str) -> str:
    """
    Handles common ZoomInfo person-page title shapes, e.g.:
      'John Doe - Director of IT - CompanyName | ZoomInfo'
      'John Doe Email & Phone Number | CompanyName Director of IT Contact Information'
      "John Doe's Email & Phone Number | CompanyName"
      'John Doe - CompanyName - ZoomInfo'
      'Contact John Doe, Email - CompanyName'
      'Who is John Doe - Director of IT at CompanyName'

    Unrecognized shapes are caught by the slug cross-check in
    search_employee_names, since ZoomInfo slugs are always the name.
    """
    if not title:
        return ""

    t = title.strip().replace("’", "'")

    # Strip trailing branding
    t = re.sub(r"\s*[-|·:]\s*ZoomInfo(?:\.com)?\s*$", "", t, flags=re.IGNORECASE)

    # Strip leading "Contact " / "Who is "
    t = re.sub(r"^\s*(?:Contact|Who\s+is)\s+", "", t, flags=re.IGNORECASE)

    # Split on dash/pipe/dot separators and take first segment as candidate name
    t = re.split(r"\s+[-–—|·]\s+", t, maxsplit=1)[0].strip()

    # Strip "Email & Phone Number" / "'s Contact Information" / ", Email" tails
    t = re.sub(
        r"(?:'s)?[\s,]+(?:Email|Phone|Contact)\b.*$",
        "",
        t,
        flags=re.IGNORECASE,
    )

    # Collapse whitespace
    t = re.sub(r"\s{2,}", " ", t).strip()

    return t

def get_profile_result_url(result: dict, profile_url_re: re.Pattern) -> str:
    """
    Extract a profile URL matching profile_url_re from a SerpAPI result.

    Matches any subdomain, so country-specific LinkedIn hosts such as
    uk.linkedin.com or in.linkedin.com are supported.
    """

    # First check the normal result URL.
    link = unquote((result.get("link") or "").strip())

    match = profile_url_re.search(link)
    if match:
        return match.group(0)

    # SerpAPI/Google may expose the actual destination inside
    # about_page_link instead.
    about_link = unquote(
        (result.get("about_page_link") or "").strip()
    )

    match = profile_url_re.search(about_link)
    if match:
        return match.group(0)

    # Neither field was a direct profile URL. SerpAPI/Google
    # sometimes returns an indirect /goto? redirect instead
    # (the same quirk file search results have) — resolve it
    # before giving up.
    resolved = resolve_serpapi_result_url(link)

    match = profile_url_re.search(resolved)
    if match:
        return match.group(0)

    return ""

def parse_first_last(full_name: str) -> tuple[str, str]:
    """
    Parse a LinkedIn/Google display name into first and last name.

    Handles:
    - Honorifics: Dr., Mr., Prof., etc.
    - Generational/name suffixes: Jr., Sr., II, III, etc.
    - Comma-delimited professional credentials: CIC, CPIA, CPA, etc.
    - Middle names/initials: ignored
    - Last-name particles: de, de la, van, von, etc.
    """
    if not full_name:
        return ("", "")

    s = full_name.strip()
    s = s.replace("\u00A0", " ")
    s = re.sub(r"\s{2,}", " ", s).strip()

    # When the leading comma segment already looks like a full name,
    # everything after the first comma is credentials or a job title, e.g.
    # "Jane Doe, CISSP, CISM" / "Sarah Kim, CISSP, Security Lead" -> keep
    # only the name. Single-token leads ("DOE, JANE") are left alone.
    if "," in s:
        segments = [seg.strip() for seg in s.split(",") if seg.strip()]

        if segments and len(segments[0].split()) >= 2:
            s = segments[0]

    # Any remaining commas are treated as separators.
    s = re.sub(r",+", " ", s)
    s = re.sub(r"\s{2,}", " ", s).strip()

    raw_parts = [p for p in s.split(" ") if p]
    parts = [normalize_name_token(p) for p in raw_parts]
    parts = [p for p in parts if p]

    if len(parts) < 2:
        return ("", "")

    # Remove leading honorifics
    while parts and parts[0].rstrip(".").lower() in HONORIFICS:
        parts.pop(0)

    if len(parts) < 2:
        return ("", "")

    # Remove trailing suffixes and credentials ("John Smith CPA")
    while parts and parts[-1].rstrip(".").lower() in TRAILING_NAME_TOKENS:
        parts.pop()

    if len(parts) < 2:
        return ("", "")

    first = parts[0]
    last = parts[-1]

    # Attach surname particles immediately preceding the last name.
    #
    # Maria De La Cruz -> first=Maria, last=De La Cruz
    i = len(parts) - 2
    particle_chain = []

    while i >= 1:
        token = parts[i].rstrip(".").lower()

        if token in LASTNAME_PARTICLES:
            particle_chain.insert(0, parts[i])
            i -= 1
            continue

        break

    if particle_chain:
        last = " ".join(particle_chain + [last])

    return (first, last)

def parse_linkedin_slug(link: str) -> tuple[str, str]:
    try:
        path = urlparse(link).path.strip("/")
        parts = path.split("/")

        if len(parts) < 2 or parts[0].lower() != "in":
            return ("", "")

        slug = unquote(parts[1]).strip().lower()

        tokens = [t for t in slug.split("-") if t]

        # Remove common LinkedIn identifier suffixes such as:
        # john-smith-8a631730
        while tokens and re.fullmatch(r"[a-f0-9]{6,}", tokens[-1]):
            tokens.pop()

        if len(tokens) < 2:
            return ("", "")

        first = normalize_name_token(tokens[0])
        last = normalize_name_token(tokens[-1])

        return (first, last)

    except Exception:
        return ("", "")

def parse_zoominfo_slug(link: str) -> tuple[str, str]:
    """
    Parse /p/John-Smith/1234567890 into ("John", "Smith"). The slug is
    run through parse_first_last so particles and suffixes are handled.
    """
    try:
        parts = urlparse(link).path.strip("/").split("/")

        if len(parts) < 2 or parts[0].lower() != "p":
            return ("", "")

        tokens = [
            t for t in unquote(parts[1]).split("-")
            if t and not t.isdigit()
        ]

        return parse_first_last(" ".join(tokens))

    except Exception:
        return ("", "")


def company_pattern(company: str) -> re.Pattern:
    """
    Regex matching the company name with any spacing/punctuation between
    its characters, so "FRSecure" also matches "FR Secure" and "AT&T"
    matches "ATT".
    """
    chars = normalize_for_email(company)
    return re.compile(r"[\W_]*".join(re.escape(c) for c in chars), re.IGNORECASE)


# Text following a company mention that marks it as part of a
# "similar people" sidebar entry rather than the person's own history:
#   "Jane Roe. Analyst. Company. Phone. Email."
#   "Company · JR. United States flag for contact location."
_sidebar_after_re = re.compile(
    r"^\s*(?:\.\s*)?(?:Phone|Email)\.|^\s*·\s*[A-Z]{2}\b|^.{0,60}flag for contact location",
    re.IGNORECASE | re.DOTALL,
)

# "Company. 2016-2019." / "Company. 2024-." job-history entries
_history_after_re = re.compile(r"^\s*\.\s*(?:19|20)\d{2}\s*-")

# Prose mention such as "role at Company", "establishing Company"
_prose_before_re = re.compile(r"\b[a-z]+[\s,]*$")


def mentions_company_as_employer(title: str, snippet: str, company_re: re.Pattern) -> bool:
    """
    True when a result ties the person to the company, as a current or
    former employee: the title names it ("... at Company"), or the snippet
    lists it in a job-history entry or mentions it in prose about the
    person. Mentions that belong to sidebar entries for other people are
    ignored.
    """
    if company_re.search(title):
        return True

    for m in company_re.finditer(snippet):
        before = snippet[max(0, m.start() - 30):m.start()]
        after = snippet[m.end():m.end() + 80]

        if _history_after_re.match(after):
            return True

        if _sidebar_after_re.match(after):
            continue

        if _prose_before_re.search(before):
            return True

    return False


@dataclass(frozen=True)
class EmployeeSource:
    key: str
    label: str
    query: str                       # Google query; {company} is filled in
    profile_url_re: re.Pattern       # identifies person/profile result URLs
    title_to_name: Callable[[str], str]
    slug_to_name: Callable[[str], tuple[str, str]]
    non_person_re: re.Pattern        # rejects directory/company titles
    slug_is_name: bool = False       # URL slug is always the real name; prefer it on mismatch
    require_company_match: bool = False  # title/snippet must tie the person to the company


EMPLOYEE_SOURCES = {
    "linkedin": EmployeeSource(
        key="linkedin",
        label="LinkedIn",
        query='site:linkedin.com/in/ "{company}"',
        profile_url_re=_linkedin_profile_url_re,
        title_to_name=clean_linkedin_title_to_name,
        slug_to_name=parse_linkedin_slug,
        non_person_re=re.compile(r"\b(linkedin|profiles?|people)\b", re.IGNORECASE),
    ),
    "zoominfo": EmployeeSource(
        key="zoominfo",
        label="ZoomInfo",
        query='site:zoominfo.com/p/ "{company}"',
        profile_url_re=_zoominfo_profile_url_re,
        title_to_name=clean_zoominfo_title_to_name,
        slug_to_name=parse_zoominfo_slug,
        non_person_re=re.compile(
            r"\b(zoominfo|profiles?|people|employees?|directory|overview|company)\b",
            re.IGNORECASE,
        ),
        slug_is_name=True,
        require_company_match=True,
    ),
}

def normalize_for_email(s: str) -> str:
    s = strip_accents(s).lower()
    s = s.replace(" ", "")
    s = s.replace("'", "")
    s = s.replace("-", "")
    s = s.replace(".", "")
    s = re.sub(r"[^a-z0-9]", "", s)
    return s


def render_email(fmt: str, first: str, last: str, email_domain: str | None = None) -> str:
    first_n = normalize_for_email(first)
    last_n = normalize_for_email(last)

    mapping = {
        "first": first_n,
        "last": last_n,
        "f": first_n[:1],
        "l": last_n[:1],
    }

    out = fmt
    for k, v in mapping.items():
        out = out.replace("{" + k + "}", v)

    if "@" not in out:
        if not email_domain:
            raise ValueError("email format does not contain '@' and no email domain was provided")
        out = out + "@" + email_domain.strip()

    return out


def search_employee_names(
    source: EmployeeSource,
    company: str,
    api_key: str,
    max_results: int,
    sleep_s: float,
    company_filter: bool = True,
) -> list[tuple[str, str, str, str, str, str]]:
    """
    Page through Google results for one employee source and return
    (source_key, profile_url, google_title, snippet, first, last) tuples.

    Raises SearchInterrupted (carrying the tuples collected so far)
    on Ctrl-C so callers can skip any remaining sources.
    """

    tag = f"({source.key})"
    query = source.query.format(company=company)

    urls_seen = set()
    results_out = []

    raw_results = 0
    profile_results = 0
    off_company = 0
    company_re = company_pattern(company)

    consecutive_empty = 0
    max_consecutive_empty = 10
    interrupted = False

    info(f"{tag} Searching {source.label}: {query}")

    try:
        for start in range(0, max_results, 10):
            info(f"{tag} Fetching results from offset {start}")

            params = {
                "engine": "google",
                "q": query,
                "api_key": api_key,
                "start": start,
                "num": 10,
            }

            results = serp_search_with_retry(
                params,
                context=f"{tag} offset {start}:",
            )

            if results is None:
                warning(
                    f"{tag} Stopping search and preserving "
                    "results collected so far."
                )
                break

            organic = results.get("organic_results", [])

            if not organic:
                consecutive_empty += 1

                notice(
                    f"{tag} Empty result page at offset {start} "
                    f"({consecutive_empty}/{max_consecutive_empty})."
                )

                if consecutive_empty >= max_consecutive_empty:
                    notice(
                        f"{tag} No results across "
                        f"{max_consecutive_empty} consecutive pages. "
                        f"Stopping."
                    )
                    break

                time.sleep(sleep_s)
                continue

            # A populated page resets the empty-page counter.
            consecutive_empty = 0

            for r in organic:
                raw_results += 1

                link = get_profile_result_url(r, source.profile_url_re)
                title = (r.get("title") or "").strip()
                snippet = (r.get("snippet") or "").strip()

                if not link:
                    continue

                profile_results += 1

                if link in urls_seen:
                    continue

                urls_seen.add(link)

                # Drop results that don't tie the person to the company, e.g.
                # ZoomInfo pages for other companies' staff that only list
                # the company in a "similar people" sidebar.
                if (
                    company_filter
                    and source.require_company_match
                    and not mentions_company_as_employer(title, snippet, company_re)
                ):
                    off_company += 1
                    continue

                name_chunk = source.title_to_name(title)

                # Ignore obvious non-person titles; the profile URL
                # itself is a person page, so the slug can still be used.
                if name_chunk and source.non_person_re.search(name_chunk):
                    name_chunk = ""

                first, last = parse_first_last(name_chunk)

                slug_first, slug_last = source.slug_to_name(link)

                # Fallback to the URL slug if Google title parsing failed,
                # or, for sources whose slug is always the person's name,
                # if the title-derived name disagrees with it.
                if slug_first and slug_last and (
                    not first
                    or not last
                    or (
                        source.slug_is_name
                        and (normalize_for_email(first), normalize_for_email(last))
                        != (normalize_for_email(slug_first), normalize_for_email(slug_last))
                    )
                ):
                    first = slug_first
                    last = slug_last

                if not first or not last:
                    continue

                # Reject abbreviated surnames such as "Tony S." / "Lisa S."
                if len(normalize_for_email(last)) < 2:
                    continue

                results_out.append((source.key, link, title, snippet, first, last))

            time.sleep(sleep_s)
    except KeyboardInterrupt:
        print()
        warning(
            f"{tag} Search interrupted by user. "
            "Preserving results collected so far."
        )
        interrupted = True

    notice(f"{tag} Raw Google results: {raw_results}")
    notice(f"{tag} {source.label} profile results: {profile_results}")
    if off_company:
        notice(
            f"{tag} Skipped {off_company} results not tied to "
            f"'{company}' (disable with --no-company-filter)."
        )
    success(f"{tag} Parsed {len(results_out)} {source.label} name hits (first+last).")

    if interrupted:
        raise SearchInterrupted(results_out)

    return results_out


# ---------------- File Enumeration Helpers ----------------
def safe_filename_from_url(url: str) -> str:
    path = urlparse(url).path
    name = os.path.basename(path) or "unknown"
    name = unquote(name)
    name = re.sub(r"[^\w.\-() ]+", "_", name).strip()
    return name or "unknown"


def guess_ext(url: str) -> str:
    parsed = urlparse(url)

    # Check query parameters first (NetSuite style)
    qs = parse_qs(parsed.query)

    if "_xt" in qs:
        ext = qs["_xt"][0]
        return ext.lower().lstrip(".")

    # Fallback to filename parsing
    fn = safe_filename_from_url(url)
    _, ext = os.path.splitext(fn)
    return ext.lower().lstrip(".")
    
PLATFORM_PATTERNS = [
    ("NetSuite", [
        r"/core/media/media\.nl\b",
        r"[?&]_xt=\.",
        r"[?&]c=\d+",
    ]),
    ("SharePoint/OneDrive", [
        r"sharepoint\.com",
        r"sharepoint-df\.com",
        r"-my\.sharepoint\.com",
        r"/_layouts/15/download\.aspx",
        r"/_layouts/15/Doc\.aspx",
        r"/:b:/s/",
        r"/personal/",
    ]),
    ("AWS S3/CDN", [
        r"\.s3\.amazonaws\.com",
        r"s3\.amazonaws\.com",
        r"\.cloudfront\.net",
    ]),
    ("Google Drive/Docs", [
        r"drive\.google\.com/file/d/",
        r"docs\.google\.com/document/d/",
        r"docs\.google\.com/spreadsheets/d/",
        r"docs\.google\.com/presentation/d/",
    ]),
    ("Salesforce", [
        r"content\.force\.com",
        r"/servlet/servlet\.FileDownload",
        r"/file-asset/",
        r"\.my\.salesforce\.com",
    ]),
    ("Thrillshare/Apptegy", [
        r"files-backend\.assets\.thrillshare\.com",
        r"\b5il\.co\b",
        r"apptegy",
    ]),
    ("Finalsite", [
        r"finalsite\.net",
        r"fs\.resource",
        r"resources\.finalsite\.net",
    ]),
    ("Edlio", [
        r"edlio\.com",
        r"edl\.io",
    ]),
    ("CivicPlus", [
        r"civicplus",
        r"civicclerk",
    ]),
]

def normalize_dt(dt) -> str:
    if not dt:
        return ""
    try:
        return dt.isoformat()
    except Exception:
        return str(dt)


def sha256_bytes(b: bytes) -> str:
    h = hashlib.sha256()
    h.update(b)
    return h.hexdigest()


def extract_internal_paths(meta_dict: dict) -> str:
    if not meta_dict:
        return ""
    hits = []
    for k, v in meta_dict.items():
        if v is None:
            continue
        s = str(v)
        if _internal_path_re.search(s):
            hits.append(f"{k}={s}")
    return "; ".join(hits)


def detect_text_encoding(sample: bytes) -> str:
    try:
        sample.decode("utf-8")
        return "utf-8"
    except Exception:
        return "latin-1"


def findings_from_text(content: bytes, sample_limit: int = 300_000) -> dict:
    sample = content[:sample_limit]
    enc = detect_text_encoding(sample)
    text = sample.decode(enc, errors="replace")

    emails = _email_re.findall(text)
    users = [u for u in _user_re.findall(text) if len(u) <= 64]
    paths = _internal_path_re.findall(text)

    kw_hits = []
    lower = text.lower()
    for kw in KEYWORDS:
        if kw.lower() in lower:
            kw_hits.append(kw)

    def summarize(items, max_samples=5):
        seen = set()
        uniq = []

        for i in items:
            if i in seen:
                continue

            seen.add(i)

            if len(uniq) < max_samples:
                uniq.append(i)

        return len(seen), uniq

    email_count, email_samples = summarize(emails)
    user_count, user_samples = summarize(users)
    path_count, path_samples = summarize(paths)

    findings_parts = []
    if email_count:
        findings_parts.append(f"emails={email_count} samples={email_samples}")
    if user_count:
        findings_parts.append(f"user_tokens={user_count} samples={user_samples}")
    if path_count:
        findings_parts.append(f"internal_paths={path_count} samples={path_samples}")
    if kw_hits:
        findings_parts.append(f"keywords={sorted(set(kw_hits))}")

    return {
        "Encoding": enc,
        "Findings": "; ".join(findings_parts) if findings_parts else "",
        "InternalPathIndicators": extract_internal_paths({"ContentSample": text[:5000]}),
    }


@dataclass
class MetaRow:
    URL: str
    FileType: str
    FileName: str
    Platform: str
    SizeBytes: str
    ContentType: str
    SHA256: str

    Title: str
    Author: str
    Creator: str
    Producer: str
    Application: str
    Company: str
    LastModifiedBy: str
    Created: str
    Modified: str

    HttpLastModified: str
    HttpETag: str

    Encoding: str
    Findings: str

    InternalPathIndicators: str
    Error: str


def http_fetch(url: str, timeout: int, max_bytes: int, user_agent: str, session: requests.Session):
    headers = {"User-Agent": user_agent}
    with session.get(url, headers=headers, timeout=timeout, stream=True, allow_redirects=True) as r:
        r.raise_for_status()

        ct = r.headers.get("Content-Type", "").split(";")[0].strip()
        lm = r.headers.get("Last-Modified", "") or ""
        etag = r.headers.get("ETag", "") or ""

        cl = r.headers.get("Content-Length")
        if cl and cl.isdigit() and int(cl) > max_bytes:
            raise ValueError(f"Content-Length {cl} exceeds max_bytes {max_bytes}")

        buf = BytesIO()
        total = 0
        for chunk in r.iter_content(chunk_size=64 * 1024):
            if not chunk:
                continue
            total += len(chunk)
            if total > max_bytes:
                raise ValueError(f"Downloaded bytes exceeded max_bytes {max_bytes}")
            buf.write(chunk)

        content = buf.getvalue()
        return content, ct, str(total), lm, etag
        
def discover_document_pages(domain, user_agent, session, timeout=5, max_workers=10):
    discovered = set()
    timed_out = 0
    unavailable = 0
    total_paths = len(DOCUMENT_PATHS)
    completed = 0
    progress_lock = threading.Lock()

    def check(path):
        url = f"https://{domain}{path}"

        try:
            r = session.get(
                url,
                headers={"User-Agent": user_agent},
                timeout=timeout,
                allow_redirects=True
            )

            if r.status_code == 200:
                return ("ok", r.url)
            return ("unavailable", None)

        except requests.exceptions.Timeout:
            return ("timeout", None)

        except requests.exceptions.RequestException:
            return ("unavailable", None)

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = [executor.submit(check, path) for path in DOCUMENT_PATHS]

        for future in as_completed(futures):
            status, url = future.result()

            with progress_lock:
                completed += 1

                if status == "ok":
                    discovered.add(url)
                elif status == "timeout":
                    timed_out += 1
                else:
                    unavailable += 1

                # Update progress on a single terminal line
                print(
                    f"\r{CYAN}[+]{RESET} Checking common document paths... "
                    f"[{completed}/{total_paths}] | Found: {len(discovered)}",
                    end="",
                    flush=True
                )

    print()

    if timed_out:
        warning(f"{timed_out} document path checks timed out.")

    if unavailable:
        notice(f"{unavailable} document paths were unavailable.")

    return discovered
    
def extract_document_links(html, base_url):
    soup = BeautifulSoup(html, "html.parser")
    urls = set()

    for tag in soup.find_all(["a", "iframe", "link"]):
        link = tag.get("href") or tag.get("src")

        if not link:
            continue

        absolute_url = urljoin(base_url, link)

        if re.search(
            r"\.(pdf|docx?|xlsx?|pptx?|csv|txt|zip)(?:\?|#|$)",
            absolute_url,
            re.I
        ):
            urls.add(absolute_url)

    return urls

def extract_json_file_urls(html):
    urls = set()

    matches = re.findall(
        r'https?://[^"\'<>\s]+?\.(?:pdf|docx?|xlsx?|pptx?|csv|txt|zip)(?:\?[^"\'<>\s]*)?',
        html,
        re.I
    )

    urls.update(matches)
    return urls
    
def compute_root_prefixes(start_pages) -> set[str]:
    """
    Derive the document-root path prefixes actually discovered for this
    domain (e.g. "/resources", "/documents", "/board") from the resolved
    start-page URLs, so folder recursion isn't hardcoded to "/documents/"
    alone.

    A start page that resolves to "/" (e.g. a dead path redirecting to
    the homepage) is excluded, since treating "/" as a root prefix would
    match every link on the site.
    """
    prefixes = set()

    for page in start_pages:
        path = urlparse(page).path.rstrip("/").lower()
        if path:
            prefixes.add(path)

    return prefixes


def is_document_folder_link(url: str, target_domain: str, root_prefixes: set[str]) -> bool:
    parsed = urlparse(url)
    host = parsed.netloc.lower()
    path = parsed.path.lower()

    target_domain = target_domain.lower().lstrip(".")

    return (
        (
            host == target_domain
            or host.endswith("." + target_domain)
        )
        and any(path.startswith(root + "/") for root in root_prefixes)
        and not re.search(
            r"\.(pdf|docx?|xlsx?|pptx?|csv|txt|zip)(?:\?|#|$)",
            path,
            re.I
        )
    )

def extract_document_folder_links(html, base_url, domain, root_prefixes):
    folder_urls = set()

    if not root_prefixes:
        return folder_urls

    # Normal href-based extraction
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup.find_all(["a"]):
        link = tag.get("href")
        if not link:
            continue

        absolute_url = urljoin(base_url, link)

        if is_document_folder_link(absolute_url, domain, root_prefixes):
            folder_urls.add(absolute_url)

    # Raw HTML / JS extraction for Apptegy/Thrillshare-style routes,
    # generalized to whichever document-ish roots were actually
    # discovered for this domain (not just "/documents/").
    root_alternation = "|".join(re.escape(root) for root in root_prefixes)
    raw_pattern = re.compile(
        r'["\']?(' + root_alternation + r'/[^"\'<>\s]+?/\d+)["\']?',
        re.I
    )

    raw_patterns = raw_pattern.findall(html)

    for path in raw_patterns:
        absolute_url = urljoin(base_url, path)

        if is_document_folder_link(absolute_url, domain, root_prefixes):
            folder_urls.add(absolute_url)

    # Handle escaped slashes from JSON: \/documents\/parents\/supply-lists\/24538815
    unescaped_html = html.replace("\\/", "/")

    escaped_patterns = raw_pattern.findall(unescaped_html)

    for path in escaped_patterns:
        absolute_url = urljoin(base_url, path)

        if is_document_folder_link(absolute_url, domain, root_prefixes):
            folder_urls.add(absolute_url)

    return folder_urls
    
def crawl_document_tree(start_pages, domain, user_agent, timeout, max_depth=5, sleep_s=0.25):
    found_files = set()
    visited_pages = set()
    queue = [(page, 0) for page in start_pages]

    root_prefixes = compute_root_prefixes(start_pages)
    if root_prefixes:
        notice(f"Recursing under discovered document root(s): {sorted(root_prefixes)}")

    try:
        while queue:
            page, depth = queue.pop(0)

            if page in visited_pages:
                continue

            if depth > max_depth:
                continue

            visited_pages.add(page)
            info(f"Crawling document folder depth={depth}: {page}")

            try:
                r = requests.get(
                    page,
                    headers={"User-Agent": user_agent},
                    timeout=timeout,
                    allow_redirects=True
                )

                if r.status_code != 200:
                    continue

                html = r.text

                found_files.update(extract_document_links(html, page))
                found_files.update(extract_json_file_urls(html))

                folder_links = extract_document_folder_links(
                    html=html,
                    base_url=page,
                    domain=domain,
                    root_prefixes=root_prefixes
                )

                for folder_url in folder_links:
                    if folder_url not in visited_pages:
                        queue.append((folder_url, depth + 1))

                time.sleep(sleep_s)

            except Exception as e:
                error(f"Failed crawling document folder {page}: {e}")
    except KeyboardInterrupt:
        raise SearchInterrupted(found_files)

    success(f"Document tree crawl visited {len(visited_pages)} pages.")
    success(f"Document tree crawl found {len(found_files)} file URLs.")

    return found_files
    
THRILLSHARE_DOMAINS = [
    "files-backend.assets.thrillshare.com",
    "5il.co",
]

def detect_platform(url: str) -> str:
    u = url.lower()

    if "thrillshare" in u:
        return "Thrillshare"

    if "5il.co" in u:
        return "Thrillshare"

    for platform, patterns in PLATFORM_PATTERNS:
        for pat in patterns:
            if re.search(pat, u, re.IGNORECASE):
                return platform

    return ""

CONTENT_TYPE_MAP = {
    "application/pdf": "pdf",
    "application/msword": "doc",
    "application/vnd.ms-excel": "xls",
    "application/vnd.ms-powerpoint": "ppt",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "text/csv": "csv",
    "text/plain": "txt",
    "application/zip": "zip",
}

# ---------------- Extractors ----------------
def extract_pdf(content: bytes) -> dict:
    if PdfReader is None:
        raise RuntimeError("pypdf not installed (python3 -m pip install pypdf)")
    reader = PdfReader(BytesIO(content))
    meta = reader.metadata
    md = {}
    if meta:
        md = {str(k): "" if v is None else str(v) for k, v in dict(meta).items()}

    return {
        "Title": md.get("/Title", ""),
        "Author": md.get("/Author", ""),
        "Creator": md.get("/Creator", ""),
        "Producer": md.get("/Producer", ""),
        "Application": "",
        "Company": "",
        "LastModifiedBy": "",
        "Created": md.get("/CreationDate", ""),
        "Modified": md.get("/ModDate", ""),
        "InternalPathIndicators": extract_internal_paths(md),
    }


def extract_docx(content: bytes) -> dict:
    if Document is None:
        raise RuntimeError("python-docx not installed (python3 -m pip install python-docx)")
    doc = Document(BytesIO(content))
    cp = doc.core_properties
    md = {
        "Title": cp.title or "",
        "Author": cp.author or "",
        "Creator": "",
        "Producer": "",
        "Application": getattr(cp, "application", "") or "",
        "Company": getattr(cp, "company", "") or "",
        "LastModifiedBy": cp.last_modified_by or "",
        "Created": normalize_dt(cp.created),
        "Modified": normalize_dt(cp.modified),
    }
    md["InternalPathIndicators"] = extract_internal_paths(md)
    return md


def extract_xlsx(content: bytes) -> dict:
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed (python3 -m pip install openpyxl)")
    wb = openpyxl.load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    p = wb.properties
    md = {
        "Title": p.title or "",
        "Author": p.creator or "",
        "Creator": p.creator or "",
        "Producer": "",
        "Application": getattr(p, "application", "") or "",
        "Company": getattr(p, "company", "") or "",
        "LastModifiedBy": p.lastModifiedBy or "",
        "Created": normalize_dt(p.created),
        "Modified": normalize_dt(p.modified),
    }
    md["InternalPathIndicators"] = extract_internal_paths(md)
    return md


def extract_pptx(content: bytes) -> dict:
    if Presentation is None:
        raise RuntimeError("python-pptx not installed (python3 -m pip install python-pptx)")
    pres = Presentation(BytesIO(content))
    cp = pres.core_properties
    md = {
        "Title": cp.title or "",
        "Author": cp.author or "",
        "Creator": "",
        "Producer": "",
        "Application": getattr(cp, "application", "") or "",
        "Company": getattr(cp, "company", "") or "",
        "LastModifiedBy": cp.last_modified_by or "",
        "Created": normalize_dt(cp.created),
        "Modified": normalize_dt(cp.modified),
    }
    md["InternalPathIndicators"] = extract_internal_paths(md)
    return md


def extract_ole_office(content: bytes) -> dict:
    if olefile is None:
        raise RuntimeError("olefile not installed (python3 -m pip install olefile)")

    ole = olefile.OleFileIO(BytesIO(content))
    meta = olefile.OleMetadata()
    meta.parse(ole)
    ole.close()

    md = {
        "Title": meta.title or "",
        "Author": meta.author or "",
        "Creator": "",
        "Producer": "",
        "Application": meta.creating_application or "",
        "Company": meta.company or "",
        "LastModifiedBy": meta.last_saved_by or "",
        "Created": normalize_dt(getattr(meta, "create_time", None)),
        "Modified": normalize_dt(getattr(meta, "last_saved_time", None)),
    }
    md["InternalPathIndicators"] = extract_internal_paths(md)
    return md


def extract_txt(content: bytes) -> dict:
    return findings_from_text(content)


def extract_csv(content: bytes) -> dict:
    return findings_from_text(content)
    
def extract_zip(content: bytes) -> dict:

    try:

        with zipfile.ZipFile(BytesIO(content)) as z:

            names = z.namelist()

        return {
            "Findings": f"zip_entries={len(names)}",
            "InternalPathIndicators":
                extract_internal_paths(
                    {"Entries": "; ".join(names[:100])}
                ),
        }

    except Exception as e:

        return {
            "Findings": "",
            "InternalPathIndicators": "",
            "Error": str(e),
        }


EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "xlsx": extract_xlsx,
    "pptx": extract_pptx,
    "doc": extract_ole_office,
    "xls": extract_ole_office,
    # opt-in
    "txt": extract_txt,
    "csv": extract_csv,
    "zip": extract_zip,
}

def resolve_serpapi_result_url(link: str, timeout: int = 10) -> str:
    """
    Resolve indirect Google /goto URLs returned by SerpAPI.

    Direct HTTP(S) result URLs are returned unchanged.
    """

    if not link:
        return ""

    link = link.strip()

    # Normal direct result
    if link.startswith(("http://", "https://")):
        return link

    # Google indirect result
    if link.startswith("/goto?"):
        google_url = "https://www.google.com" + link

        try:
            r = requests.get(
                google_url,
                headers={"User-Agent": "Mozilla/5.0"},
                allow_redirects=True,
                timeout=timeout,
                stream=True,
            )

            return r.url

        except requests.RequestException:
            return ""

    return ""

def serp_search_filetype(domain: str, ext: str, api_key: str, max_results: int, sleep_s: float) -> set[str]:
    q = f"site:{domain} filetype:{ext}"
    urls = set()

    try:
        _serp_search_filetype_pages(q, domain, ext, api_key, max_results, sleep_s, urls)
    except KeyboardInterrupt:
        raise SearchInterrupted(urls)

    return urls


def _serp_search_filetype_pages(q, domain, ext, api_key, max_results, sleep_s, urls: set[str]):
    for start in range(0, max_results, 10):
        info(f"({ext}) Fetching results from offset {start}")
        params = {"engine": "google", "q": q, "api_key": api_key, "start": start, "num": 10}
        
        results = serp_search_with_retry(
            params,
            context=f"({ext}) offset {start}:",
        )

        if results is None:
            warning(
                f"({ext}) Stopping search and preserving "
                f"results collected so far."
            )
            break
        
        organic = results.get("organic_results", [])
        
        if not organic:
            notice(f"({ext}) No more results.")
            break

        for r in organic:
            raw_link = (r.get("link") or "").strip()

            if not raw_link:
                continue

            link = resolve_serpapi_result_url(raw_link)

            if not link:
                continue

            parsed = urlparse(link)
            hostname = parsed.netloc.lower()

            domain_lower = domain.lower().lstrip(".")

            allowed_domain = (
                hostname == domain_lower
                or hostname.endswith("." + domain_lower)
            )

            allowed = (
                allowed_domain
                or bool(detect_platform(link))
            )

            if allowed:
                urls.add(link)

        time.sleep(sleep_s)


def build_metadata_row(url: str, args, session: requests.Session) -> MetaRow:
    row = MetaRow(
        URL=url,
        FileType="unknown",
        FileName="",
        Platform="",
        SizeBytes="",
        ContentType="",
        SHA256="",

        Title="",
        Author="",
        Creator="",
        Producer="",
        Application="",
        Company="",
        LastModifiedBy="",
        Created="",
        Modified="",

        HttpLastModified="",
        HttpETag="",

        Encoding="",
        Findings="",

        InternalPathIndicators="",
        Error="",
    )

    try:
        ext = guess_ext(url) or "unknown"
        row.FileType = ext
        row.FileName = safe_filename_from_url(url)
        row.Platform = detect_platform(url)

        content, ct, size_bytes, lm, etag = http_fetch(
            url=url, timeout=args.timeout, max_bytes=args.max_bytes,
            user_agent=args.user_agent, session=session
        )
        ct = ct.split(";")[0].strip()
        row.ContentType = ct
        # If extension unknown, try to infer from Content-Type
        if ext not in EXTRACTORS and ct in CONTENT_TYPE_MAP:
            ext = CONTENT_TYPE_MAP[ct]
            row.FileType = ext
        row.SizeBytes = size_bytes
        row.HttpLastModified = lm
        row.HttpETag = etag
        row.SHA256 = sha256_bytes(content)

        extractor = EXTRACTORS.get(ext)
        if not extractor:
            row.Error = f"No extractor for extension: {ext}"
        else:
            md = extractor(content)
            for k, v in md.items():
                if hasattr(row, k) and v is not None:
                    setattr(row, k, str(v))

    except Exception as e:
        row.Error = str(e)

    return row


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Enumerate publicly indexed files via SerpAPI (Google) and extract "
            "high-value metadata."
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
    examples:
      File enumeration:
        python3 SirPapiSearch.py example.com

      LinkedIn email enumeration:
        python3 SirPapiSearch.py example.com --linkedin --company "Example Company" --email-format "{f}{last}"

      LinkedIn email enumeration, multiple candidate formats:
        python3 SirPapiSearch.py example.com --linkedin --company "Example Company" --email-format "{f}{last},{first}.{last},{first}{l}"

      ZoomInfo email enumeration:
        python3 SirPapiSearch.py example.com --zoominfo --company "Example Company" --email-format "{f}{last}"

      Employee enumeration across all sources (LinkedIn + ZoomInfo):
        python3 SirPapiSearch.py example.com --employees --company "Example Company" --email-format "{f}{last}"

      Employee enumeration, selected sources only:
        python3 SirPapiSearch.py example.com --employees --sources zoominfo --company "Example Company" --email-format "{f}{last}"
    """,
    )
    parser.add_argument("domain", help="Target domain for file enumeration OR email domain for employee modes (e.g. example.com)")

    parser.add_argument(
        "--api-key",
        default=None,
        help="SerpAPI key (overrides SERPAPI_KEY env var and HARDCODED_SERPAPI_KEY)"
    )
    
    parser.add_argument(
        "--no-banner",
        action="store_true",
        help="Suppress the SirPapiSearch banner"
    )

    # Employee modes: OFF by default. When enabled, tool will ONLY run
    # employee enumeration and exit.
    employee_modes = parser.add_mutually_exclusive_group()
    employee_modes.add_argument(
        "--linkedin",
        action="store_true",
        help="LinkedIn email enumeration mode (SerpAPI Google results only). Requires --company and --email-format. Exits after writing emails."
    )
    employee_modes.add_argument(
        "--zoominfo",
        action="store_true",
        help="ZoomInfo email enumeration mode (SerpAPI Google results only). Requires --company and --email-format. Exits after writing emails."
    )
    employee_modes.add_argument(
        "--employees",
        action="store_true",
        help="Employee email enumeration across multiple sources (default: all; see --sources). "
             "Requires --company and --email-format. Exits after writing emails."
    )
    parser.add_argument(
        "--sources",
        default=None,
        help="Comma-separated sources for --employees "
             f"(choices: {','.join(EMPLOYEE_SOURCES)}; default: all). "
             "--max applies per source."
    )
    parser.add_argument("--company", default=None,
                        help="Company name to search for (required with --linkedin/--zoominfo/--employees)")
    parser.add_argument(
        "--no-company-filter",
        action="store_true",
        help="Keep ZoomInfo results whose title/snippet don't tie the person to --company "
             "(by default they are skipped as likely other companies' staff)"
    )
    parser.add_argument("--email-format", default=None,
                        help="REQUIRED with --linkedin/--zoominfo/--employees. Template supports {first},{last},{f},{l}. "
                             "Accepts a comma-separated list to generate multiple candidate "
                             "formats per contact in one pass, e.g. "
                             "'{f}{last},{first}.{last},{first}{l}'. "
                             "Examples: '{f}{last}@domain.com', '{first}{last}@domain.com', '{first}.{last}' (appends @<domain>).")
    # Optional override; by default we use the positional domain argument as the email domain.
    parser.add_argument("--email-domain", default=None,
                        help="Optional override: email domain to use for employee modes. If omitted, uses positional <domain> argument.")
    parser.add_argument("--out-emails", default=None,
                        help="Output file for generated emails (default: <mode>-emails.txt, "
                             "where <mode> is linkedin, zoominfo or employees)")

    parser.add_argument(
        "--out-profiles",
        default=None,
        help="Output CSV for source profiles (default: <mode>-profiles.csv)"
    )

    parser.add_argument(
    "--out-names",
    default=None,
    help=(
        "Output file for discovered first/last name pairs "
        "(default: <mode>-names.txt)."
    ),
    )

    parser.add_argument(
        "--types",
        default="pdf,docx,xlsx,pptx,doc,xls",
        help="Comma-separated file extensions (default: pdf,docx,xlsx,pptx,doc,xls). Add csv,txt if desired."
    )
    parser.add_argument("--max", type=int, default=400, help="Max SerpAPI results per type (default: 400)")
    parser.add_argument("--sleep", type=float, default=0.5, help="Sleep between SerpAPI requests (default: 0.5)")
    parser.add_argument("--timeout", type=int, default=20, help="HTTP timeout seconds (default: 20)")
    parser.add_argument("--max-bytes", type=int, default=20_000_000, help="Max download size per file (default: 20MB)")
    parser.add_argument("--workers", type=int, default=10, help="Concurrent worker threads for file downloads/probing (default: 10)")
    parser.add_argument("--user-agent", default="Mozilla/5.0 (compatible; FileEnum/3.1)",
                        help="User-Agent for HTTP fetches")
    parser.add_argument("--out-urls", default=None, help="Output file for URLs (default: <domain>-URLs.txt)")
    parser.add_argument("--out-csv", default=None, help="Output CSV file (default: <domain>-Metadata.csv)")
    args = parser.parse_args()

    if not args.no_banner:
        print_banner()

    # Resolve API key priority:
    api_key = args.api_key or os.getenv("SERPAPI_KEY") or HARDCODED_SERPAPI_KEY
    if not api_key:
        raise SystemExit(
            "[-] Missing SerpAPI key. Provide --api-key, set SERPAPI_KEY env var, "
            "or hardcode HARDCODED_SERPAPI_KEY in the script."
        )
        
    if args.sources and not args.employees:
        raise SystemExit("[-] --sources can only be used with --employees")

    # ---------------- Employee Modes (only if explicitly requested) ----------------
    if args.linkedin or args.zoominfo or args.employees:
        if args.linkedin:
            mode, source_keys = "linkedin", ["linkedin"]
        elif args.zoominfo:
            mode, source_keys = "zoominfo", ["zoominfo"]
        else:
            mode = "employees"
            source_keys = [
                s.strip().lower()
                for s in (args.sources or ",".join(EMPLOYEE_SOURCES)).split(",")
                if s.strip()
            ]
            # Keep order, drop duplicates
            source_keys = list(dict.fromkeys(source_keys))

            unknown = [s for s in source_keys if s not in EMPLOYEE_SOURCES]
            if unknown or not source_keys:
                raise SystemExit(
                    f"[-] Unknown --sources value(s): {', '.join(unknown) or '(none)'}. "
                    f"Choices: {', '.join(EMPLOYEE_SOURCES)}"
                )

        if not args.company:
            raise SystemExit(f"[-] --company is required when using --{mode}")
        if not args.email_format:
            raise SystemExit(f"[-] --email-format is required when using --{mode}")

        out_emails = args.out_emails or f"{mode}-emails.txt"
        out_profiles = args.out_profiles or f"{mode}-profiles.csv"
        out_names = args.out_names or f"{mode}-names.txt"

        # Comma-separated list of candidate formats, e.g.
        # "{f}{last},{first}.{last},{first}{l}"
        email_formats = [
            fmt.strip()
            for fmt in args.email_format.split(",")
            if fmt.strip()
        ]

        if not email_formats:
            raise SystemExit("[-] --email-format did not contain any usable formats")

        # Email domain defaults to positional <domain> unless overridden
        effective_email_domain = args.email_domain or args.domain
        
        if "." not in effective_email_domain:
            raise SystemExit(
                f"[-] Email domain looks invalid: '{effective_email_domain}'. "
                f"Use a full domain like '{effective_email_domain}.com' or pass --email-domain."
            )
            
        contacts = []

        try:
            for key in source_keys:
                contacts.extend(
                    search_employee_names(
                        source=EMPLOYEE_SOURCES[key],
                        company=args.company,
                        api_key=api_key,
                        max_results=args.max,
                        sleep_s=args.sleep,
                        company_filter=not args.no_company_filter,
                    )
                )
        except KeyboardInterrupt as e:
            contacts.extend(getattr(e, "partial_results", []))
            if len(source_keys) > 1:
                warning("Skipping remaining sources.")

        if len(source_keys) > 1:
            success(
                f"({mode}) Parsed {len(contacts)} name hits across "
                f"{len(source_keys)} sources: {', '.join(source_keys)}."
            )

        emails = set()
        for (source, url, title, snippet, first, last) in contacts:
            for fmt in email_formats:
                try:
                    emails.add(
                        render_email(
                            fmt,
                            first,
                            last,
                            effective_email_domain
                        )
                    )
                except Exception as e:
                    error(
                        f"({source}) Failed rendering email for "
                        f"{first} {last} with format '{fmt}' ({url}): {e}"
                    )

        sorted_emails = sorted(emails)
        # write <mode>-emails.txt
        with open(out_emails, "w", encoding="utf-8") as f:
            for e in sorted_emails:
                f.write(e + "\n")

        success(
            f"Emails saved to {out_emails} "
            f"({len(sorted_emails)} unique across {len(email_formats)} format(s))."
        )

        # Write profile/source information for auditing
        with open(out_profiles, "w", newline="", encoding="utf-8") as csvfile:
            writer = csv.writer(csvfile)

            writer.writerow([
                "Source",
                "ProfileURL",
                "GoogleTitle",
                "FirstName",
                "LastName",
                "Snippet"
            ])

            for source, url, title, snippet, first, last in contacts:
                writer.writerow([
                    source,
                    url,
                    title,
                    first,
                    last,
                    snippet
                ])

        success(
        f"Profiles saved to {out_profiles} "
        f"({len(contacts)} profiles)."
        )

        # Write unique first/last name pairs for later email mangling.
        # Case-insensitive so the same person found by several sources
        # is only listed once.
        names = {}
        for _, _, _, _, first, last in contacts:
            names.setdefault(f"{first} {last}".lower(), f"{first} {last}")

        sorted_names = sorted(names.values(), key=str.lower)

        with open(out_names, "w", encoding="utf-8") as f:
            for name in sorted_names:
                f.write(name + "\n")

        success(
            f"Names saved to {out_names} "
            f"({len(sorted_names)} unique names)."
        )

        return  # will not proceed automatically with file enumeration afterward in employee modes
        
    if BeautifulSoup is None:
        warning(
            "beautifulsoup4 not installed. "
            "Document portal crawling disabled. "
            "Install with: python3 -m pip install beautifulsoup4"
        )

    # ---------------- File Enumeration Mode (default) ----------------
    out_csv = args.out_csv or f"{args.domain}-Metadata.csv"
    out_urls = args.out_urls or f"{args.domain}-URLs.txt"
    types = [t.strip().lower().lstrip(".") for t in args.types.split(",") if t.strip()]
    all_urls: set[str] = set()

    session = requests.Session()

    try:
        for ext in types:
            all_urls |= serp_search_filetype(args.domain, ext, api_key, args.max, args.sleep)

        if BeautifulSoup is not None:

            document_pages = discover_document_pages(
                args.domain,
                args.user_agent,
                session,
                timeout=5,
                max_workers=args.workers
            )

            all_urls.update(
                crawl_document_tree(
                    start_pages=document_pages,
                    domain=args.domain,
                    user_agent=args.user_agent,
                    timeout=args.timeout,
                    max_depth=5,
                    sleep_s=0.25
                )
            )
    except KeyboardInterrupt as e:
        print()
        all_urls |= getattr(e, "partial_results", set())
        warning(
            "URL collection interrupted by user. Skipping remaining "
            "searches and processing the URLs collected so far "
            "(Ctrl-C again to stop metadata extraction)."
        )

    sorted_urls = sorted(all_urls)
    print()
    success(f"Found {len(sorted_urls)} unique URLs across types: {', '.join(types)}")

    with open(out_urls, "w", encoding="utf-8") as f:
        for u in sorted_urls:
            f.write(u + "\n")
    success(f"URLs saved to {out_urls}")

    rows_by_url = {}
    completed = 0

    executor = ThreadPoolExecutor(max_workers=args.workers)
    try:
        futures = {
            executor.submit(build_metadata_row, url, args, session): url
            for url in sorted_urls
        }

        for future in as_completed(futures):
            url = futures[future]
            rows_by_url[url] = future.result()

            completed += 1
            info(f"({completed}/{len(sorted_urls)}) Processed: {url}")

        executor.shutdown()
    except KeyboardInterrupt:
        print()
        warning(
            "Metadata extraction interrupted by user. "
            f"Saving the {len(rows_by_url)}/{len(sorted_urls)} rows processed so far."
        )
        # Drop queued downloads; in-flight ones finish in the background.
        executor.shutdown(wait=False, cancel_futures=True)

    fieldnames = list(MetaRow.__annotations__.keys())
    with open(out_csv, "w", newline="", encoding="utf-8") as csvfile:
        writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
        writer.writeheader()

        for url in sorted_urls:
            if url in rows_by_url:
                writer.writerow(asdict(rows_by_url[url]))

    success(f"Report saved to {out_csv}")


if __name__ == "__main__":
    main()
